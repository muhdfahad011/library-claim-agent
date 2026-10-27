import json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

data = json.load(open("valued.json", encoding="utf-8"))

wb = Workbook()
ws = wb.active
ws.title = "Inventory"

ws["A1"] = "Home Library Contents Claim Packet"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"Generated {date.today().isoformat()} | Prices are model estimates, not verified retail quotes | ISBNs are representative editions"

headers = ["#", "Title", "Author", "Publisher", "First published",
           "ISBN-13 (representative)", "Format", "Est. replacement value (USD)",
           "Price confidence", "Price source", "Status", "Review notes"]
ws.append([])
ws.append(headers)
hdr_row = 4
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=hdr_row, column=c)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="305496")
    cell.alignment = Alignment(wrap_text=True, vertical="center")

review_fill = PatternFill("solid", fgColor="FFF2CC")
for i, b in enumerate(data, start=1):
    m = b.get("match") or {}
    v = b.get("valuation") or {}
    ws.append([
        i, b["title"], b.get("author"), m.get("publisher"),
        m.get("first_published"), m.get("isbn13"), v.get("format"),
        v.get("price_usd"), v.get("confidence"), v.get("source"),
        b.get("status"), "; ".join(b.get("review_reasons", [])),
    ])
    row = ws.max_row
    ws.cell(row=row, column=8).number_format = '"$"#,##0.00'
    if b.get("status") == "REVIEW":
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).fill = review_fill

first, last = hdr_row + 1, ws.max_row
tot = last + 2
ws.cell(row=tot, column=7, value="TOTAL").font = Font(bold=True)
t = ws.cell(row=tot, column=8, value=f"=SUM(H{first}:H{last})")
t.font = Font(bold=True)
t.number_format = '"$"#,##0.00'
ws.cell(row=tot + 1, column=7, value="Items").font = Font(bold=True)
ws.cell(row=tot + 1, column=8, value=f"=COUNTA(B{first}:B{last})")
ws.cell(row=tot + 2, column=7, value="Need review").font = Font(bold=True)
ws.cell(row=tot + 2, column=8, value=f'=COUNTIF(K{first}:K{last},"REVIEW")')

widths = [5, 42, 24, 24, 14, 22, 12, 20, 12, 16, 10, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"

wb.save("claim_packet.xlsx")
total = sum((b.get("valuation") or {}).get("price_usd") or 0 for b in data)
print(f"saved claim_packet.xlsx | {len(data)} items | total ${total:,.2f}")
